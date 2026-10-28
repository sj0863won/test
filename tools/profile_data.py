#!/usr/bin/env python3
"""Generic data profiler for an unfamiliar task folder (Python standard library only).

Usage:  python test/tools/profile_data.py [folder=.] [--out notes/profile.md] [--max-rows 200000]

Reads files only (never modifies them) and writes a short Markdown summary:
  - tables (csv/tsv/json/jsonl/xlsx): rows, columns, fill rate, value types, ranges,
    date formats and time-zone marks, duplicates, odd-format values (typo candidates),
    values that differ only by case/spaces, memo-like columns
  - key candidates between files and values found on one side only
  - documents/logs: line count, timestamp formats, lines with warning words
    (revision, effective date, notice, exception, UTC, duplicate, retest ...)
This is a first look for humans and the coding agent. Always confirm in the raw files.
"""
import argparse
import csv
import io
import json
import math
import os
import re
import sys
import time
from collections import Counter, defaultdict

SKIP_DIRS = {'.git', 'node_modules', '__pycache__', '.venv', 'venv', '.idea', '.vscode'}
ROOT_SKIP = {'test', 'harness', 'notes', 'output', 'outputs'}   # top level only: the harness itself, notes and results (data/test/ etc. are still read)
TABLE_EXT = {'.csv', '.tsv'}
JSON_EXT = {'.json', '.jsonl', '.ndjson'}
TEXT_EXT = {'.txt', '.md', '.log', '.eml', '.yaml', '.yml', '.ini', '.cfg', '.toml', '.sql', '.xml', '.html', '.htm'}
EMPTY = {'', 'na', 'n/a', 'nan', 'null', 'none', '-', '--', '?', '#n/a', 'nil'}
WARN_WORDS = re.compile(
    r'(개정|rev\.?\s?\d|버전|version|\bv\d+(\.\d+)+|시행|적용일|발효|공지|예외|주의|금지|단위|UTC|KST|GMT|시간대|timezone'
    r'|중복|재테스트|재측정|재시험|retest|정정|오타|누락|협의|임시|초안|draft|deprecated|폐기|우선|기준일|미정|TBD|TODO'
    r'|변경|교체|정비|교정|calibrat|override|무시|제외|단, )', re.I)
MEMO_COL = re.compile(r'(memo|note|비고|메모|comment|remark|reason|사유|특이|remarks)', re.I)
IDISH_COL = re.compile(r'(^id$|_id$|^id_|serial|시리얼|번호|\bno\b|_no$|code|코드|key|lot|tool|설비|장비|chamber|챔버|name|이름|item|항목)', re.I)
DATEISH_COL = re.compile(r'(date|time|일자|일시|시각|날짜|_at$|^at$|ts$|timestamp|day)', re.I)

NUM_RE = re.compile(r'^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$')
NUM_COMMA_RE = re.compile(r'^[+-]?\d{1,3}(,\d{3})+(\.\d+)?$')
NUM_UNIT_RE = re.compile(r'^[+-]?\d+(\.\d+)?\s*(%|‰|°C|℃|°|ppm|dppm|[a-zA-Zμµ]+(/[a-zA-Z]+)?)$')
DATE_RES = [
    ('YYYY-MM-DD', re.compile(r'^\d{4}-\d{2}-\d{2}$')),
    ('YYYY-MM-DD hh:mm', re.compile(r'^\d{4}-\d{2}-\d{2} \d{1,2}:\d{2}(:\d{2}(\.\d+)?)?$')),
    ('YYYY-MM-DDThh:mm', re.compile(r'^\d{4}-\d{2}-\d{2}T\d{1,2}:\d{2}(:\d{2}(\.\d+)?)?$')),
    ('...Z (UTC)', re.compile(r'^\d{4}-\d{2}-\d{2}[T ]\d{1,2}:\d{2}(:\d{2}(\.\d+)?)?Z$')),
    ('...+hh:mm (offset)', re.compile(r'^\d{4}-\d{2}-\d{2}[T ]\d{1,2}:\d{2}(:\d{2}(\.\d+)?)?\s?[+-]\d{2}:?\d{2}$')),
    ('... KST/UTC 글자', re.compile(r'^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}[T ]\d{1,2}:\d{2}(:\d{2})?\s?(KST|UTC|GMT)$', re.I)),
    ('YYYY/MM/DD·YYYY.MM.DD', re.compile(r'^\d{4}[/.]\d{1,2}[/.]\d{1,2}\.?([ T]\d{1,2}:\d{2}(:\d{2})?)?$')),
    ('MM/DD/YYYY', re.compile(r'^\d{1,2}/\d{1,2}/\d{4}([ T]\d{1,2}:\d{2}(:\d{2})?)?$')),
    ('DD-Mon-YYYY', re.compile(r'^\d{1,2}[- ][A-Za-z]{3}[- ]\d{2,4}( \d{1,2}:\d{2}(:\d{2})?)?$')),
    ('hh:mm', re.compile(r'^\d{1,2}:\d{2}(:\d{2})?$')),
]
YYYYMMDD = re.compile(r'^(19|20)\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])$')
TS_IN_TEXT = re.compile(r'(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}[T ]\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?)(Z|\s?[+-]\d{2}:?\d{2}|\s?(?:KST|UTC|GMT))?')
DATE_IN_TEXT = re.compile(r'\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b')
MAX_DISTINCT = 20000


def short(v, n=40):
    v = str(v).replace('\n', ' ').replace('|', '/')
    return v if len(v) <= n else v[:n - 1] + '…'


def detect_encoding(path):
    with open(path, 'rb') as f:
        head = f.read(1 << 20)
    full = len(head) < (1 << 20)
    for enc in ('utf-8-sig', 'cp949'):
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError as e:
            if not full and e.start >= len(head) - 4:
                try:
                    head[:e.start].decode(enc)
                    return enc
                except UnicodeDecodeError:
                    pass
    return 'latin-1'


def shape(v):
    out = []
    for ch in v:
        if ch.isdigit():
            out.append('9')
        elif 'A' <= ch <= 'Z':
            out.append('A')
        elif 'a' <= ch <= 'z':
            out.append('a')
        elif '가' <= ch <= '힣':
            out.append('가')
        elif ch.isspace():
            out.append('_')
        else:
            out.append(ch)
    s = ''.join(out)
    s = re.sub(r'(A|a|가)\1+', lambda m: m.group(1) + '+', s)
    s = re.sub(r'9{7,}', '9+', s)
    s = re.sub(r'_{2,}', '_', s)
    return s[:40]


def date_kind(v, colname):
    for name, rx in DATE_RES:
        if rx.match(v):
            return name
    if YYYYMMDD.match(v) and DATEISH_COL.search(colname or ''):
        return 'YYYYMMDD'
    return None


class Col:
    def __init__(self, name):
        self.name = name
        self.n = 0
        self.empty = 0
        self.empty_marks = Counter()
        self.values = Counter()
        self.capped = False
        self.num = []
        self.non_num_examples = Counter()
        self.comma_num = 0
        self.unit_num = Counter()
        self.date_kinds = Counter()
        self.date_min = {}
        self.date_max = {}
        self.ws = 0
        self.shapes = Counter()
        self.shape_ex = {}
        self.lens = 0
        self.nested = 0
        self.nonfinite = Counter()

    def add(self, raw):
        self.n += 1
        if raw is None:
            self.empty += 1
            return
        if isinstance(raw, (list, dict)):
            self.nested += 1
            raw = '[%s len %d]' % (type(raw).__name__, len(raw))
        v = str(raw)
        s = v.strip()
        if s.lower() in EMPTY:
            self.empty += 1
            if s:
                self.empty_marks[s] += 1
            return
        if s != v:
            self.ws += 1
        self.lens += len(s)
        if s in self.values or len(self.values) < MAX_DISTINCT:
            self.values[s] += 1
        else:
            self.capped = True
        fv = float(s) if NUM_RE.match(s) else None
        if fv is not None and not math.isfinite(fv):
            self.nonfinite[s] += 1          # e.g. 1e309 -> inf: report it, never use it in number stats
            fv = None
        if fv is not None:
            if len(self.num) < 300000:
                self.num.append(fv)
        else:
            if NUM_COMMA_RE.match(s):
                self.comma_num += 1
            m = NUM_UNIT_RE.match(s)
            if m:
                self.unit_num[m.group(2)] += 1
            if len(self.non_num_examples) < 50 or s in self.non_num_examples:
                self.non_num_examples[s] += 1
        dk = date_kind(s, self.name)
        if dk:
            self.date_kinds[dk] += 1
            if dk not in self.date_min or s < self.date_min[dk]:
                self.date_min[dk] = s
            if dk not in self.date_max or s > self.date_max[dk]:
                self.date_max[dk] = s
        sh = shape(s)
        self.shapes[sh] += 1
        self.shape_ex.setdefault(sh, s)

    @property
    def filled(self):
        return self.n - self.empty

    def kind(self):
        f = self.filled
        if f == 0:
            return 'empty'
        if self.nested and self.nested >= 0.5 * f:
            return 'nested'
        if sum(self.date_kinds.values()) >= 0.8 * f:
            return 'date'
        if len(self.num) >= 0.95 * f:
            return 'number'
        if len(self.num) >= 0.6 * f:
            return 'number?'
        distinct = len(self.values)
        avg = self.lens / max(f, 1)
        if distinct <= 20:
            return 'category'
        if avg > 40 or len(self.shapes) > 12:
            return 'text'
        return 'id'

    def summary_and_flags(self):
        k = self.kind()
        f = self.filled
        flags = []
        summ = ''
        if self.empty_marks:
            flags.append('빈 값 표기: ' + ', '.join("'%s' %d" % (short(a, 8), c) for a, c in self.empty_marks.most_common(3)))
        if self.ws:
            flags.append('앞뒤 공백 %d' % self.ws)
        if k == 'nested':
            summ = '중첩 값(list/dict) %d' % self.nested
        elif k == 'date':
            parts = []
            for dk, c in self.date_kinds.most_common():
                parts.append('%s ~ %s [%s %d]' % (self.date_min[dk], self.date_max[dk], dk, c))
            summ = ' / '.join(parts[:2])
            if len(self.date_kinds) > 1:
                flags.append('날짜 형식 %d종 혼재' % len(self.date_kinds))
            if any(dk.startswith('...Z') or 'offset' in dk or 'KST' in dk for dk in self.date_kinds):
                flags.append('시간대 표기 있음(%s)' % ', '.join(dk for dk in self.date_kinds if dk.startswith('...')))
            nd = f - sum(self.date_kinds.values())
            if nd:
                flags.append('날짜 아닌 값 %d (예: %s)' % (nd, ', '.join(short(x, 16) for x, _ in self.non_num_examples.most_common(2) if not date_kind(x, self.name))[:60]))
        elif k in ('number', 'number?'):
            xs = sorted(self.num)
            mid = xs[len(xs) // 2]
            isint = all(x == int(x) for x in xs[:2000])
            fmt = (lambda x: '%d' % x) if isint else (lambda x: '%.4g' % x)
            summ = '%s ~ %s (중앙 %s)' % (fmt(xs[0]), fmt(xs[-1]), fmt(mid))
            zeros = sum(1 for x in xs if x == 0)
            negs = sum(1 for x in xs if x < 0)
            if zeros:
                summ += ' · 0이 %d' % zeros
            if negs:
                summ += ' · 음수 %d' % negs
            nn = f - len(self.num)
            if nn:
                ex = ', '.join("'%s'" % short(x, 12) for x, _ in self.non_num_examples.most_common(3))
                flags.append('숫자 아닌 값 %d (예: %s)' % (nn, ex))
        elif k == 'category':
            summ = ', '.join('%s %d' % (short(v, 18), c) for v, c in self.values.most_common(12))
            if len(self.values) > 12:
                summ += ' …'
        elif k == 'text':
            ex = [v for v, _ in self.values.most_common(2)]
            summ = '자유 텍스트 (평균 %d자) 예: %s' % (self.lens // max(f, 1), ' / '.join(short(x, 30) for x in ex))
        else:  # id-like
            distinct = len(self.values)
            summ = '고유 %s · 예: %s' % (('%d+' % distinct) if self.capped else distinct, ', '.join(short(v, 22) for v, _ in self.values.most_common(3)))
        if k in ('id', 'number?') and len(self.shapes) > 1 and not MEMO_COL.search(self.name):
            total = sum(self.shapes.values())
            top_sh, top_c = self.shapes.most_common(1)[0]
            if top_c >= 0.75 * total:
                rare = [(sh, c) for sh, c in self.shapes.most_common() if sh != top_sh]
                rare_n = sum(c for _, c in rare)
                if rare_n:
                    flags.append('형식이 튀는 값 %d (예: %s) ← 오타·표기 확인' % (rare_n, ', '.join(short(self.shape_ex[sh], 22) for sh, _ in rare[:4])))
            elif len(self.shapes) <= 6:
                flags.append('표기 혼재: ' + ', '.join('%s %d%%' % (short(self.shape_ex[sh], 16), round(100 * c / total)) for sh, c in self.shapes.most_common(4)))
        if k in ('id', 'category'):
            groups = defaultdict(set)
            for v in self.values:
                groups[re.sub(r'[\s_\-]+', '', v.lower())].add(v)
            variants = [sorted(g) for g in groups.values() if len(g) > 1]
            if variants:
                flags.append('대소문자·공백·구분자만 다른 값 %d묶음 (예: %s)' % (len(variants), ' | '.join('/'.join(short(x, 14) for x in g[:3]) for g in variants[:3])))
        if self.nonfinite:
            flags.append('표현 범위를 넘는 숫자 %d (예: %s) ← 입력 오류 확인' % (sum(self.nonfinite.values()), ', '.join(short(x, 12) for x in self.nonfinite)[:40]))
        if self.comma_num:
            flags.append('천 단위 쉼표 숫자 %d' % self.comma_num)
        if self.unit_num:
            flags.append('단위가 붙은 숫자 %d (%s)' % (sum(self.unit_num.values()), ', '.join(self.unit_num.keys())[:30]))
        return k, summ, flags


def read_table(path, ext, max_rows):
    enc = detect_encoding(path)
    info = {'encoding': enc}
    with open(path, 'r', encoding=enc, errors='replace', newline='') as fh:
        sample = fh.read(65536)
        fh.seek(0)
        delim = '\t' if ext == '.tsv' else ','
        if ext != '.tsv':
            try:
                delim = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
            except csv.Error:
                pass
        info['delimiter'] = {',': 'comma', '\t': 'tab', ';': 'semicolon', '|': 'pipe'}.get(delim, delim)
        reader = csv.reader(fh, delimiter=delim)
        try:
            header = next(reader)
        except StopIteration:
            return info, [], []
        header = [h.strip() or ('col%d' % (i + 1)) for i, h in enumerate(header)]
        rows = []
        bad = 0
        total = 0
        for r in reader:
            total += 1
            if len(r) != len(header):
                bad += 1
            if len(rows) < max_rows:
                rows.append(r)
        info['rows_total'] = total
        info['ragged'] = bad
    return info, header, rows


def read_json(path, ext, max_rows):
    enc = detect_encoding(path)
    info = {'encoding': enc}
    recs = []
    if ext in ('.jsonl', '.ndjson'):
        bad = 0
        total = 0
        with open(path, 'r', encoding=enc, errors='replace') as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                total += 1
                try:
                    o = json.loads(line)
                except ValueError:
                    bad += 1
                    continue
                if len(recs) < max_rows:
                    recs.append(o if isinstance(o, dict) else {'value': o})
        info['rows_total'] = total
        info['bad_lines'] = bad
    else:
        with open(path, 'r', encoding=enc, errors='replace') as fh:
            try:
                o = json.load(fh)
            except ValueError as e:
                info['error'] = 'JSON 파싱 실패: %s' % e
                return info, [], []
        if isinstance(o, dict):
            lists = [(k, v) for k, v in o.items() if isinstance(v, list) and v and isinstance(v[0], dict)]
            info['top_keys'] = list(o.keys())[:15]
            if lists:
                k, v = max(lists, key=lambda kv: len(kv[1]))
                info['records_key'] = k
                o = v
            else:
                return info, [], []
        if isinstance(o, list):
            recs = [x if isinstance(x, dict) else {'value': x} for x in o[:max_rows]]
            info['rows_total'] = len(o)
    header = []
    seen = set()
    for r in recs[:5000]:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                header.append(k)
    rows = [[r.get(h) for h in header] for r in recs]
    return info, header, rows


def read_xlsx(path, max_rows):
    """Return one (sheet_name, info, header, rows) per sheet. Every sheet is profiled, not only the first."""
    try:
        import openpyxl  # noqa
    except ImportError:
        return [('', {'error': 'openpyxl 없음 - xlsx는 직접 확인'}, [], [])]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for name in wb.sheetnames:
        it = wb[name].iter_rows(values_only=True)
        try:
            header = [str(h) if h is not None else 'col%d' % (i + 1) for i, h in enumerate(next(it))]
        except StopIteration:
            out.append((name, {'error': '빈 시트', 'sheets': wb.sheetnames}, [], []))
            continue
        rows, total = [], 0
        for r in it:
            total += 1
            if len(rows) < max_rows:
                rows.append(['' if x is None else x for x in r])
        out.append((name, {'sheets': wb.sheetnames, 'rows_total': total, 'encoding': 'xlsx'}, header, rows))
    return out


def profile_text(path):
    enc = detect_encoding(path)
    lines = []
    with open(path, 'r', encoding=enc, errors='replace') as fh:
        for i, line in enumerate(fh):
            if i < 200000:
                lines.append(line.rstrip('\n'))
    first = [l for l in lines if l.strip()][:3]
    tz = Counter()
    ts = []
    dates = []
    for l in lines[:200000]:
        for m in TS_IN_TEXT.finditer(l):
            ts.append(m.group(1))
            mark = (m.group(2) or '').strip()
            tz[mark if mark else '표기 없음'] += 1
        for m in DATE_IN_TEXT.finditer(l):
            dates.append('%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3))))
    warn = []
    for i, l in enumerate(lines):
        if WARN_WORDS.search(l) and l.strip():
            warn.append('L%d: %s' % (i + 1, short(l.strip(), 140)))
    return {'encoding': enc, 'lines': len(lines), 'first': first, 'tz': tz, 'ts_n': len(ts),
            'dates': (min(dates), max(dates)) if dates else None, 'warn': warn}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('folder', nargs='?', default='.')
    ap.add_argument('--out', default=None, help='write Markdown here (default: print)')
    ap.add_argument('--max-rows', type=int, default=200000)
    a = ap.parse_args()
    t0 = time.time()
    root = os.path.abspath(a.folder)
    files = []
    for d, dirs, fs in os.walk(root):
        dirs[:] = sorted(x for x in dirs if x not in SKIP_DIRS and not x.startswith('.')
                         and not (d == root and x in ROOT_SKIP))
        for f in sorted(fs):
            if f.startswith('.'):
                continue
            files.append(os.path.join(d, f))
    out = []
    w = out.append
    tables = []
    docs = []
    others = []
    hot = []
    keyidx = defaultdict(set)      # value -> {(file, col)}
    colvals = {}
    colnames = defaultdict(list)   # normalized name -> [(file, col)]
    def register(rel, size, info, header, rows):
        cols = [Col(h) for h in header]
        dup = Counter()
        for r in rows:
            key = tuple(str(x) for x in r)
            dup[key] += 1
            for i, c in enumerate(cols):
                c.add(r[i] if i < len(r) else None)
        dup_rows = sum(c - 1 for c in dup.values() if c > 1)
        tables.append((rel, size, info, header, cols, dup_rows, len(rows)))
        for c in cols:
            if c.name == '_file':
                continue
            k = c.kind()
            nm = re.sub(r'[\s_\-]+', '', c.name.lower())
            colnames[nm].append((rel, c.name))
            if k in ('id', 'category', 'number') and len(c.values) >= 3:
                if k == 'number' and not IDISH_COL.search(c.name):
                    continue
                vals = set(v for v in c.values if len(v) >= 2)
                if k == 'number':
                    vals = set(v for v in vals if len(v) >= 3)
                colvals[(rel, c.name)] = vals
                for v in vals:
                    keyidx[v].add((rel, c.name))

    # many small JSON files in one folder (one record per file) -> one table (+ one table per nested list)
    jdirs = defaultdict(list)
    for p in files:
        if p.lower().endswith('.json'):
            jdirs[os.path.dirname(p)].append(p)
    grouped_json = set()
    for d, ps in jdirs.items():
        if len(ps) < 5:
            continue
        recs, nested = [], defaultdict(list)
        for p in ps:
            try:
                with open(p, 'r', encoding=detect_encoding(p), errors='replace') as fh:
                    o = json.load(fh)
            except Exception:
                continue
            if not isinstance(o, dict):
                continue
            name = os.path.basename(p)
            row = {'_file': name}
            for k, v in o.items():
                if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                    for item in v:
                        nested[k].append(dict({'_file': name}, **{kk: vv for kk, vv in item.items()}))
                    row[k] = v
                elif isinstance(v, dict):
                    for kk, vv in v.items():
                        row['%s.%s' % (k, kk)] = vv
                else:
                    row[k] = v
            recs.append(row)
            grouped_json.add(p)
        if not recs:
            continue
        reld = os.path.relpath(d, root).replace(os.sep, '/')
        for label, items in [('', recs)] + [('[%s]' % k, v) for k, v in nested.items()]:
            header, seen = [], set()
            for r in items:
                for k in r:
                    if k not in seen:
                        seen.add(k)
                        header.append(k)
            register('%s/*.json%s (%d개 파일)' % (reld, label, len(recs)), 0, {'encoding': 'utf-8-sig', 'rows_total': len(items)},
                     header, [[r.get(h) for h in header] for r in items])

    for p in files:
        if p in grouped_json:
            continue
        rel = os.path.relpath(p, root).replace(os.sep, '/')
        ext = os.path.splitext(p)[1].lower()
        size = os.path.getsize(p)
        try:
            if ext in TABLE_EXT or ext in JSON_EXT or ext in ('.xlsx', '.xlsm'):
                if ext in TABLE_EXT:
                    parts = [('',) + read_table(p, ext, a.max_rows)]
                elif ext in JSON_EXT:
                    parts = [('',) + read_json(p, ext, a.max_rows)]
                else:
                    parts = read_xlsx(p, a.max_rows)
                for sheet, info, header, rows in parts:
                    name = '%s[%s]' % (rel, sheet) if sheet else rel
                    if not header:
                        docs.append((name, size, info, None))
                        continue
                    register(name, size, info, header, rows)
            elif ext in TEXT_EXT or (size < 2_000_000 and ext in ('', '.dat', '.out')):
                docs.append((rel, size, None, profile_text(p)))
            else:
                others.append((rel, size))
        except Exception as e:  # keep going on odd files
            others.append((rel, size, 'error: %s' % e))

    # key links via inverted index
    pair = Counter()
    for v, cs in keyidx.items():
        if len(cs) < 2 or len(cs) > 8:
            continue
        cs = sorted(cs)
        for i in range(len(cs)):
            for j in range(i + 1, len(cs)):
                if cs[i][0] != cs[j][0]:
                    pair[(cs[i], cs[j])] += 1
    links = []
    for (A, B), inter in pair.items():
        na, nb = len(colvals[A]), len(colvals[B])
        ratio = inter / max(1, min(na, nb))
        if inter >= 3 and ratio >= 0.2:
            only_a = sorted(colvals[A] - colvals[B])
            only_b = sorted(colvals[B] - colvals[A])
            links.append((ratio * min(na, nb), A, B, inter, na, nb, only_a, only_b))
    links.sort(key=lambda x: -x[0])
    # collapse many columns of one file linking to the same target column
    grouped = defaultdict(list)
    for L in links:
        _, A, B, inter, na, nb, oa, ob = L
        grouped[('L', A[0], B)].append(L)      # many columns of A's file -> one column B
        grouped[('R', B[0], A)].append(L)      # many columns of B's file -> one column A
    collapsed = []
    used = set()
    for L in links:
        _, A, B, *_ = L
        gl, gr = grouped[('L', A[0], B)], grouped[('R', B[0], A)]
        if len(gl) >= 4:
            key = ('L', A[0], B)
            if key not in used:
                used.add(key)
                collapsed.append(('group', A[0], [x[1][1] for x in gl], B, sum(x[3] for x in gl)))
        elif len(gr) >= 4:
            key = ('R', B[0], A)
            if key not in used:
                used.add(key)
                collapsed.append(('group', B[0], [x[2][1] for x in gr], A, sum(x[3] for x in gr)))
        else:
            collapsed.append(('pair',) + L)

    same_name_diff = []
    for nm, lst in colnames.items():
        if len(lst) < 2 or len(nm) < 2:
            continue
        for i in range(len(lst)):
            for j in range(i + 1, len(lst)):
                A, B = lst[i], lst[j]
                if A[0] == B[0] or A not in colvals or B not in colvals:
                    continue
                va, vb = colvals[A], colvals[B]
                if not va or not vb:
                    continue
                inter = len(va & vb)
                sa = Counter(shape(x) for x in va).most_common(1)[0][0]
                sb = Counter(shape(x) for x in vb).most_common(1)[0][0]
                if inter / max(1, min(len(va), len(vb))) < 0.5 and sa != sb:
                    same_name_diff.append((A, B, inter, sorted(va)[:3], sorted(vb)[:3]))

    # ---------- write ----------
    w('# 데이터 첫 훑기 (test/tools/profile_data.py)')
    w('')
    w('폴더 `%s` · 파일 %d개 · %.1f초 · 자동 요약이므로 의심 지점은 원본에서 다시 확인하세요.' % (a.folder, len(files), time.time() - t0))
    w('')
    w('## 파일 목록')
    for rel, size, info, header, cols, dup_rows, n in tables:
        tot = info.get('rows_total', n)
        extra = []
        if info.get('encoding') not in ('utf-8-sig', 'xlsx'):
            extra.append('인코딩 ' + info.get('encoding'))
            hot.append('%s: 인코딩이 %s (UTF-8 아님)' % (rel, info.get('encoding')))
        if info.get('ragged'):
            extra.append('열 수가 다른 행 %d' % info['ragged'])
            hot.append('%s: 열 수가 헤더와 다른 행 %d' % (rel, info['ragged']))
        if info.get('bad_lines'):
            extra.append('깨진 줄 %d' % info['bad_lines'])
        if dup_rows:
            extra.append('완전 중복 행 %d' % dup_rows)
            hot.append('%s: 완전히 같은 행 %d개' % (rel, dup_rows))
        if tot > n:
            extra.append('앞 %d행만 분석' % n)
        w('- 표 `%s` — %s행 × %d열%s' % (rel, format(tot, ','), len(header), (' · ' + ' · '.join(extra)) if extra else ''))
    for rel, size, info, tx in docs:
        if tx:
            w('- 문서/로그 `%s` — %s줄%s' % (rel, format(tx['lines'], ','), '' if tx['encoding'] == 'utf-8-sig' else ' · 인코딩 ' + tx['encoding']))
        else:
            w('- `%s` — 표로 읽지 못함 %s' % (rel, info.get('error', '') if info else ''))
    for o in others:
        w('- 기타 `%s` — %s bytes%s' % (o[0], format(o[1], ','), (' · ' + o[2]) if len(o) > 2 else ''))
    w('')

    w('## 파일 간 연결 후보 (같은 값이 들어 있는 열)')
    if not collapsed:
        w('- 찾지 못함 (값이 겹치는 ID성 열이 없음)')
    for c in collapsed[:40]:
        if c[0] == 'group':
            _, fa, names, B, inter = c
            w('- `%s`의 %d개 열(%s …) → `%s`.%s' % (fa, len(names), ', '.join(names[:4]), B[0], B[1]))
        else:
            _, score, A, B, inter, na, nb, oa, ob = c
            line = '- `%s`.%s ↔ `%s`.%s — 공통 %d (왼쪽 고유 %d, 오른쪽 고유 %d)' % (A[0], A[1], B[0], B[1], inter, na, nb)
            small_only = oa if na <= nb else ob
            side = '왼쪽' if na <= nb else '오른쪽'
            if small_only and len(small_only) <= max(3, 0.2 * min(na, nb)):
                line += ' · %s에만 있는 값 %d: %s' % (side, len(small_only), ', '.join(short(x, 20) for x in small_only[:5]))
            if small_only and len(small_only) <= max(3, 0.2 * min(na, nb)) and min(na, nb) >= 5:
                hot.append('%s.%s 값 중 %d개가 %s.%s에 없음 (예: %s)' % (
                    (A if na <= nb else B)[0], (A if na <= nb else B)[1], len(small_only),
                    (B if na <= nb else A)[0], (B if na <= nb else A)[1], ', '.join(short(x, 20) for x in small_only[:3])))
            w(line)
    for A, B, inter, ea, eb in same_name_diff[:10]:
        w('- 이름은 같은데 값이 거의 안 겹침: `%s`.%s (예: %s) vs `%s`.%s (예: %s) ← 표기 체계가 다른지 확인' % (
            A[0], A[1], ', '.join(short(x, 14) for x in ea), B[0], B[1], ', '.join(short(x, 14) for x in eb)))
        hot.append('같은 이름의 열 %s.%s / %s.%s 값 표기가 다름' % (A[0], A[1], B[0], B[1]))
    w('')

    w('## 표별 열 요약')
    for rel, size, info, header, cols, dup_rows, n in tables:
        w('')
        w('### `%s`' % rel)
        meta = []
        if info.get('records_key'):
            meta.append('레코드 위치: "%s"' % info['records_key'])
        if info.get('delimiter') and info['delimiter'] != 'comma':
            meta.append('구분자 ' + info['delimiter'])
        if meta:
            w('(%s)' % ' · '.join(meta))
        w('| 열 | 채움 | 고유 | 종류 | 요약 | 주의 |')
        w('|---|---|---|---|---|---|')
        for c in cols:
            try:
                k, summ, flags = c.summary_and_flags()
            except Exception as e:  # one odd column must not stop the whole report
                k, summ, flags = '?', '', ['요약 실패: %s' % e]
            fill = '%d%%' % round(100 * c.filled / max(c.n, 1))
            if c.empty:
                fill += ' (빈 %d)' % c.empty
            uniq = ('%d+' % len(c.values)) if c.capped else str(len(c.values))
            w('| %s | %s | %s | %s | %s | %s |' % (short(c.name, 24), fill, uniq, k, short(summ, 110), '; '.join(flags)))
            for fl in flags:
                if any(t in fl for t in ('형식이 튀는', '표기 혼재', '대소문자', '숫자 아닌', '날짜 형식', '시간대', '단위가 붙은', '천 단위', '빈 값 표기')):
                    hot.append('%s.%s: %s' % (rel, c.name, fl))
            if MEMO_COL.search(c.name) and k in ('category', 'text', 'id'):
                vals = [v for v, _ in c.values.most_common(8)]
                if vals:
                    w('| ↳ %s 값 | | | | %s | 비고·메모는 전부 읽어 볼 것 |' % (short(c.name, 20), short(' / '.join(vals), 200)))
        if dup_rows:
            w('- 완전 중복 행 %d개' % dup_rows)
        idc = [c for c in cols if c.kind() in ('id', 'number') and IDISH_COL.search(c.name)]
        for c in idc[:3]:
            rep = [(v, k) for v, k in c.values.items() if k > 1]
            if rep and len(c.values) > 0.3 * c.filled:
                w('- `%s` 같은 값이 2번 이상 나오는 경우 %d개 (예: %s) ← 재측정·중복·여러 행 구조인지 확인' % (
                    c.name, len(rep), ', '.join('%s×%d' % (short(v, 18), k) for v, k in sorted(rep, key=lambda x: -x[1])[:3])))

    if docs:
        w('')
        w('## 문서·로그')
        for rel, size, info, tx in docs:
            if not tx:
                continue
            w('')
            w('### `%s` (%s줄)' % (rel, format(tx['lines'], ',')))
            for l in tx['first']:
                w('    ' + short(l, 140))
            if tx['ts_n']:
                w('- 시각 %d개 · 시간대 표기: %s' % (tx['ts_n'], ', '.join('%s %d' % (k, v) for k, v in tx['tz'].most_common())))
                if len(tx['tz']) > 1:
                    hot.append('%s: 시각의 시간대 표기가 섞여 있음 (%s)' % (rel, ', '.join(tx['tz'].keys())))
                elif any(k in ('Z', 'UTC') or k.startswith('+00') for k in tx['tz']):
                    hot.append('%s: 시각이 UTC로 기록됨' % rel)
            if tx['dates']:
                w('- 날짜 범위: %s ~ %s' % tx['dates'])
            if tx['warn']:
                w('- 주의 단어가 든 줄 (%d개 중 앞 12개):' % len(tx['warn']))
                for x in tx['warn'][:12]:
                    w('  - ' + x)

    w('')
    w('## 먼저 확인할 신호 (자동 추출, 중요도 순 아님)')
    seen = set()
    k = 0
    for h in hot:
        if h in seen:
            continue
        seen.add(h)
        k += 1
        w('%d. %s' % (k, short(h, 220)))
        if k >= 40:
            break
    if not k:
        w('- 자동으로 찾은 신호 없음 (직접 확인 필요)')
    text = '\n'.join(out) + '\n'
    if a.out:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
        with open(a.out, 'w', encoding='utf-8') as fh:
            fh.write(text)
        print('wrote %s (%d lines, %.1fs)' % (a.out, len(out), time.time() - t0))
    else:
        sys.stdout.write(text)


if __name__ == '__main__':
    main()
